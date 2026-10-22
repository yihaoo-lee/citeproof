import csv

from docx import Document
from openpyxl import load_workbook

from app.models import Claim, Matrix, MatrixRow
from app.services.export import export_matrix


def test_three_formats_open_and_hide_unverified_summary(tmp_path):
    matrix = Matrix(query="test", rows=[MatrixRow(
        paper_id="p1", title="Paper", fields={"problem": "Unverified private summary", "method": "Verified method"},
        claims=[Claim(field="method", quote="Original method.", page=1, verified=True, entailment="supported")],
    )])
    paths = export_matrix(matrix, tmp_path)
    with open(paths["csv"], encoding="utf-8-sig", newline="") as handle:
        csv_row = next(csv.DictReader(handle))
    assert csv_row["研究问题"] == "未找到原文出处"
    assert csv_row["方法"] == "Verified method"
    workbook = load_workbook(paths["xlsx"], read_only=True)
    assert "未找到原文出处" in [cell.value for cell in workbook.active[2]]
    doc = Document(paths["docx"])
    content = "\n".join(p.text for p in doc.paragraphs)
    assert "未找到原文出处" in content
    assert "Unverified private summary" not in content
