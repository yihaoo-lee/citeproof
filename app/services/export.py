"""导出 CSV / XLSX / DOCX。"""
from __future__ import annotations

import csv
from pathlib import Path

from ..models import FIELD_NAMES, Matrix

FIELD_LABEL_ZH = {
    "problem": "研究问题",
    "method": "方法",
    "dataset": "数据集",
    "conclusion": "结论",
    "limitation": "局限",
    "reproducibility": "可复现性",
}
LEVEL_ZH = {"exact": "精确匹配", "normalized": "归一化匹配", "anchored": "锚点定位"}
EVIDENCE_ZH = {"fulltext": "全文证据", "abstract": "摘要级证据"}


def _display_field(row, field: str) -> str:
    value = row.fields.get(field, "")
    if not value:
        return ""
    return value if any(c.field == field and c.verified for c in row.claims) else "未找到原文出处"


def _flat_rows(matrix: Matrix) -> list[dict]:
    rows = []
    for r in matrix.rows:
        row = {
            "paper_id": r.paper_id,
            "title": r.title,
            "year": r.year,
            "language": r.language,
        }
        for f in FIELD_NAMES:
            row[FIELD_LABEL_ZH[f]] = _display_field(r, f)
        row["引用数"] = len(r.claims)
        row["引用明细"] = " | ".join(
            f"[{FIELD_LABEL_ZH.get(c.field, c.field)}] 第{c.page}页 {c.quote}" for c in r.claims
        )
        rows.append(row)
    return rows


def export_matrix(matrix: Matrix, out_dir) -> dict:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = _flat_rows(matrix)
    fieldnames = ["paper_id", "title", "year", "language"] + list(FIELD_LABEL_ZH.values()) + ["引用数", "引用明细"]

    csv_path = out_dir / "matrix.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    xlsx_path = out_dir / "matrix.xlsx"
    import pandas as pd

    pd.DataFrame(rows, columns=fieldnames).to_excel(xlsx_path, index=False)

    docx_path = out_dir / "matrix.docx"
    from docx import Document

    doc = Document()
    doc.add_heading("有据 CiteProof · 文献证据矩阵", 0)
    doc.add_paragraph(f"研究问题：{matrix.query}")
    doc.add_paragraph(f"生成时间：{matrix.generated_at}")
    stats = matrix.stats
    pct = lambda v: f"{v * 100:.1f}%"
    doc.add_paragraph(
        f"论文数 {stats.papers}｜可追溯率 {pct(stats.quote_hit_rate)}｜丢弃率 {pct(stats.discard_rate)}"
        f"｜支持率(模型判定) {pct(stats.support_rate)}｜字段完整率 {pct(stats.field_completeness)}"
    )
    for r in matrix.rows:
        doc.add_heading(r.title, level=1)
        for f in FIELD_NAMES:
            if r.fields.get(f):
                doc.add_paragraph(f"{FIELD_LABEL_ZH[f]}：{_display_field(r, f)}")
        for c in r.claims:
            doc.add_paragraph(
                f"引用（{FIELD_LABEL_ZH.get(c.field, c.field)}，第 {c.page} 页，"
                f"{LEVEL_ZH.get(c.match_level, c.match_level)}，"
                f"{EVIDENCE_ZH.get(c.evidence_level, c.evidence_level)}）：{c.quote}"
            )
    doc.save(docx_path)

    return {"csv": str(csv_path), "xlsx": str(xlsx_path), "docx": str(docx_path)}
